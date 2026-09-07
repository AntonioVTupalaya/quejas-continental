import os
import re
import openpyxl

BASE_DIR = os.path.dirname(__file__)
CARGA_LECTIVA = os.path.join(BASE_DIR, "Carga Lectiva 202620.xlsx")
DIRECTORIO_DOCENTES = os.path.join(BASE_DIR, "Directorio_Docentes_2026-20.xlsx")

# Días válidos de la semana (para el selector de horario)
DIAS_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

# Abreviaturas usadas en la columna HOR
MAPA_DIAS = {
    "LUN": "Lunes",
    "MAR": "Martes",
    "MIE": "Miércoles",
    "JUE": "Jueves",
    "VIE": "Viernes",
    "SAB": "Sábado",
    "DOM": "Domingo",
}


def leer_carga_lectiva():
    """Lee la carga lectiva y devuelve una lista de dicts con NRC, curso, docente y horarios."""
    wb = openpyxl.load_workbook(CARGA_LECTIVA, data_only=True)
    ws = wb["Carga Lectiva"]

    registros = []
    for r in range(2, ws.max_row + 1):
        nrc = ws.cell(r, 6).value
        curso = ws.cell(r, 8).value
        docente = ws.cell(r, 27).value
        hor = ws.cell(r, 28).value
        if nrc is None or curso is None:
            continue
        registros.append(
            {
                "nrc": str(nrc).strip(),
                "curso": str(curso).strip(),
                "docente": str(docente).strip().rstrip("*") if docente else "",
                "hor": str(hor) if hor else "",
            }
        )
    wb.close()
    return registros


def parsear_horario(hor):
    """Parsea la columna HOR y devuelve una lista de bloques {dia, inicio, fin}.

    Formato ejemplo: '(CLAS)JUE 2030:2159(-Sin asignar)' o con '|' para varios bloques.
    """
    if not hor:
        return []
    bloques = []
    for parte in str(hor).split("|"):
        parte = parte.strip()
        if not parte.startswith("("):
            continue
        dia_abr = parte[6:9]
        dia = MAPA_DIAS.get(dia_abr, dia_abr)
        # Buscar el patrón de horario HHHH:HHMM dentro de la parte
        coincidencia = re.search(r"(\d{4}):(\d{4})", parte)
        if not coincidencia:
            continue
        inicio = f"{coincidencia.group(1)[:2]}:{coincidencia.group(1)[2:]}"
        fin = f"{coincidencia.group(2)[:2]}:{coincidencia.group(2)[2:]}"
        bloques.append({"dia": dia, "inicio": inicio, "fin": fin})
    return bloques


def obtener_info_nrc(nrc):
    """Dado un NRC, devuelve {curso, docente, horarios} o None si no existe."""
    nrc = str(nrc).strip()
    registros = leer_carga_lectiva()
    coincidencias = [r for r in registros if r["nrc"] == nrc]
    if not coincidencias:
        return None
    primero = coincidencias[0]
    horarios = []
    for reg in coincidencias:
        horarios.extend(parsear_horario(reg["hor"]))
    return {
        "curso": primero["curso"],
        "docente": primero["docente"],
        "horarios": horarios,
    }


def leer_directorio_docentes():
    """Lee el directorio de docentes y devuelve lista de dicts."""
    wb = openpyxl.load_workbook(DIRECTORIO_DOCENTES, data_only=True)
    ws = wb["Sheet1"]
    docentes = []
    for r in range(2, ws.max_row + 1):
        nombre = ws.cell(r, 5).value
        celular = ws.cell(r, 7).value
        correo = ws.cell(r, 10).value
        if nombre:
            docentes.append(
                {
                    "nombre": str(nombre).strip(),
                    "celular": str(celular).strip() if celular else "",
                    "correo": str(correo).strip() if correo else "",
                }
            )
    wb.close()
    return docentes


def leer_horas_pdf(archivo_pdf):
    """Extrae horas de inicio y fin del PDF de horarios de la universidad."""
    import pdfplumber
    inicios = set()
    fines = set()
    with pdfplumber.open(archivo_pdf) as pdf:
        for p in pdf.pages:
            texto = p.extract_text() or ""
            for m in re.finditer(r"(\d{2}:\d{2})-(\d{2}:\d{2})", texto):
                inicios.add(m.group(1))
                fines.add(m.group(2))
    return sorted(inicios), sorted(fines)


HORARIO_PDF = os.path.join(BASE_DIR, "HORARIO PRESENCIAL 2026-20.pdf")
HORARIOS_GENERALES = os.path.join(BASE_DIR, "HORARIOS_GENERALES_2026-20.xlsx")

# Días que comparten el horario de lunes a viernes
DIAS_SEMANA_LABORAL = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]

# Franjas de horario por día: {dia: [{"inicio": "HH:MM", "fin": "HH:MM"}, ...]}
# Tomadas del archivo HORARIOS_GENERALES_2026-20.xlsx (lunes a viernes, sábado y domingo).
FRANJAS_POR_DIA = {}


def _fmt_hora(t):
    return f"{t.hour:02d}:{t.minute:02d}"


def _leer_franjas(ws, col_inicio, col_fin):
    franjas = []
    for r in range(4, ws.max_row + 1):
        vi = ws.cell(r, col_inicio).value
        vf = ws.cell(r, col_fin).value
        if vi is not None and vf is not None:
            franjas.append({"inicio": _fmt_hora(vi), "fin": _fmt_hora(vf)})
    return franjas


def _inicializar_franjas():
    global FRANJAS_POR_DIA
    if FRANJAS_POR_DIA:
        return
    import openpyxl
    wb = openpyxl.load_workbook(HORARIOS_GENERALES, data_only=True)
    ws = wb["HORARIO"]
    semana = _leer_franjas(ws, 1, 2)
    sabado = _leer_franjas(ws, 4, 5)
    domingo = _leer_franjas(ws, 7, 8)
    wb.close()
    for d in DIAS_SEMANA_LABORAL:
        FRANJAS_POR_DIA[d] = semana
    FRANJAS_POR_DIA["Sábado"] = sabado
    FRANJAS_POR_DIA["Domingo"] = domingo


_inicializar_franjas()

# Listas de horas de inicio y fin por día para los selectores del formulario
INICIOS_POR_DIA = {d: [f["inicio"] for f in franjas] for d, franjas in FRANJAS_POR_DIA.items()}
FINES_POR_DIA = {d: [f["fin"] for f in franjas] for d, franjas in FRANJAS_POR_DIA.items()}

# Listas genéricas (por si se necesitan sin distinguir día)
HORAS_INICIO = list(dict.fromkeys(i for d in INICIOS_POR_DIA.values() for i in d))
HORAS_FIN = list(dict.fromkeys(f for d in FINES_POR_DIA.values() for f in d))


def _minutos(hh):
    a, b = hh.split(":")
    return int(a) * 60 + int(b)


def calcular_fin(inicio):
    """Dado un inicio, devuelve el fin = inicio + 2 horas."""
    mins = _minutos(inicio) + 120
    h = mins // 60
    m = mins % 60
    return f"{h:02d}:{m:02d}"


if __name__ == "__main__":
    info = obtener_info_nrc("31435")
    print(info)
